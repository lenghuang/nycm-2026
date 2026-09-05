import dlt
import openpyxl

from config.settings import MacroFactorSettings

DAILY_COLUMNS = {
    "Date": "date",
    "Expenditure": "expenditure_kcal",
    "Trend Weight (lb)": "trend_weight_lbs",
    "Weight (lb)": "weight_lbs",
    "Calories (kcal)": "calories_kcal",
    "Protein (g)": "protein_g",
    "Fat (g)": "fat_g",
    "Carbs (g)": "carbs_g",
    "Target Calories (kcal)": "target_calories_kcal",
    "Target Protein (g)": "target_protein_g",
    "Target Fat (g)": "target_fat_g",
    "Target Carbs (g)": "target_carbs_g",
    "Steps": "steps",
    "Alcohol (g)": "alcohol_g",
    "Fiber (g)": "fiber_g",
    "Sodium (mg)": "sodium_mg",
    "Caffeine (mg)": "caffeine_mg",
    "Calcium (mg)": "calcium_mg",
    "Iron (mg)": "iron_mg",
    "Vitamin D (mcg)": "vitamin_d_mcg",
    "Omega-3 (g)": "omega3_g",
}

FOOD_LOG_COLUMNS = {
    "Date": "date",
    "Time": "time",
    "Food Name": "food_name",
    "Serving Size": "serving_size",
    "Serving Qty": "serving_qty",
    "Serving Weight (g)": "serving_weight_g",
    "Calories (kcal)": "calories_kcal",
    "Protein (g)": "protein_g",
    "Fat (g)": "fat_g",
    "Carbs (g)": "carbs_g",
    "Fiber (g)": "fiber_g",
    "Sodium (mg)": "sodium_mg",
}


@dlt.source(name="macrofactor")
def macrofactor_source(settings: MacroFactorSettings):
    @dlt.resource(name="daily_summary", write_disposition="merge", primary_key="date")
    def daily_summary():
        wb = openpyxl.load_workbook(str(settings.export_path))
        ws = wb["Quick Export"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            raw = dict(zip(headers, row))
            record = {clean: raw[orig] for orig, clean in DAILY_COLUMNS.items() if orig in raw}
            if hasattr(record.get("date"), "date"):
                record["date"] = record["date"].date()
            yield record

    @dlt.resource(name="food_log", write_disposition="merge", primary_key=["date", "time", "food_name"])
    def food_log():
        wb = openpyxl.load_workbook(str(settings.export_path))
        ws = wb["Food Log"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            raw = dict(zip(headers, row))
            yield {clean: raw[orig] for orig, clean in FOOD_LOG_COLUMNS.items() if orig in raw}

    yield daily_summary
    yield food_log
